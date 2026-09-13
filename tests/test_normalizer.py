"""The normalizer's converters, the email repairs and the binary-as-text
guard (wayfinder #28). ``tests/test_west_desk_fixtures.py`` pins the
fixtures themselves; this file converts them and checks the result.
"""

import re
from datetime import datetime
from pathlib import Path

import pytest

from strata.normalizer import Conversion, Refusal, decode_text, normalize

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "examples" / "west-desk" / "sources"
MAIL = SOURCES / "mail"


def _convert(relative: str) -> Conversion:
    path = SOURCES / relative
    result = normalize(path.read_bytes(), relative)
    assert isinstance(result, Conversion), f"{relative}: refused, {result}"
    return result


# --- email: header recovery, footer, quoting -------------------------------


@pytest.mark.parametrize(
    "name",
    [
        "plain",
        "quoted-angle",
        "outlook-original-message",
        "interleaved",
        "thread-parent",
        "thread-child",
        "attachment",
        "quoted-printable-wrapped",
        "display-names-only",
        "received-only",
        "empty-body",
    ],
)
def test_every_mail_fixture_converts_with_no_footer_in_any_paragraph(name):
    conversion = _convert(f"mail/{name}.eml")
    for paragraph in conversion.paragraphs:
        assert "EDRM Enron Email Data Set has been produced" not in paragraph


def test_the_first_paragraph_is_the_header_block_as_text():
    conversion = _convert("mail/display-names-only.eml")
    header = conversion.paragraphs[0]
    assert header.startswith("From: ")
    assert "Kessler, Ruth; Lindqvist, Tomas; Baptiste, Corinne; PV" in header


def test_quoted_angle_drops_every_depth_of_quoting():
    conversion = _convert("mail/quoted-angle.eml")
    body = "\n".join(conversion.paragraphs[1:])
    assert "Agreed, moving it. Marcus signed off on the call." in body
    assert ">" not in body


def test_outlook_original_message_block_is_cut_solid():
    conversion = _convert("mail/outlook-original-message.eml")
    body = "\n".join(conversion.paragraphs[1:])
    assert "Nobody schedules across the tie" in body
    assert "Original Message" not in body
    assert "Sent:" not in body
    assert "forced outage on the Cascade tie" not in body  # the quoted reply


def test_interleaved_keeps_the_senders_lines_in_order():
    conversion = _convert("mail/interleaved.eml")
    body = conversion.paragraphs[1:]
    assert body == (
        "No, it steps. Forty today, sixty on Thursday if the second line holds.",
        "Tomas is holding it. I told him to wait for your word.",
    )


def test_quoted_printable_soft_breaks_join_into_one_run():
    conversion = _convert("mail/quoted-printable-wrapped.eml")
    body = "\n".join(conversion.paragraphs)
    assert "eleven thousand megawatt hours" in body


def test_the_attachment_is_never_opened():
    conversion = _convert("mail/attachment.eml")
    for paragraph in conversion.paragraphs:
        assert "cascade-tie-load.xls" not in paragraph
    assert any("Load figures attached" in p for p in conversion.paragraphs)


def test_thread_index_decodes_to_a_shared_conversation_id_with_the_parent_index_chained():
    parent = _convert("mail/thread-parent.eml")
    child = _convert("mail/thread-child.eml")

    assert parent.metadata["thread_conversation_id"] == child.metadata["thread_conversation_id"]
    assert child.metadata["thread_parent_index"] == parent.metadata["thread_index"]


def test_empty_body_has_the_header_paragraph_only_and_an_empty_title():
    """The title comes from the body paragraphs alone (#28's "then the ref"
    fallback never sees paragraph 0, which is never empty); with no body
    words either, the title is empty here and the sources adapter supplies
    the ref."""
    conversion = _convert("mail/empty-body.eml")
    assert len(conversion.paragraphs) == 1
    assert conversion.title == ""


def test_plain_email_title_is_the_subject():
    conversion = _convert("mail/plain.eml")
    assert conversion.title == "Cascade tie congestion report"


def test_an_empty_subject_email_with_body_words_titles_from_the_body():
    raw = (
        b"From: a@example.com\r\nTo: b@example.com\r\nSubject: \r\n\r\n"
        b"The tie held through the peak.\r\n"
    )
    result = normalize(raw, "mail/x.eml")
    assert isinstance(result, Conversion)
    assert result.title == "The tie held through the peak."


# --- header paragraph normalization: unfolding, RFC 2047, raw 8-bit, Cc/Date -


def _header(raw_bytes: bytes) -> str:
    result = normalize(raw_bytes, "mail/x.eml")
    assert isinstance(result, Conversion)
    return result.paragraphs[0]


def test_a_folded_to_header_unfolds_with_no_stray_crlf_and_words_in_order():
    raw = (
        b"From: Priya <priya@example.com>\r\n"
        b"To: Ruth <ruth@example.com>,\r\n"
        b"\tTomas <tomas@example.com>,\r\n"
        b" Corinne <corinne@example.com>\r\n"
        b"Subject: Desk coverage\r\n"
        b"\r\n"
        b"Body text.\r\n"
    )
    to_line = next(line for line in _header(raw).split("\n") if line.startswith("To:"))
    assert "\r" not in to_line and "\n" not in to_line
    assert to_line.index("Ruth") < to_line.index("Tomas") < to_line.index("Corinne")


def test_a_bare_lf_folded_to_header_unfolds_with_tab_and_space_continuations():
    """The CRLF twin of test_a_folded_to_header_unfolds_with_no_stray_crlf_and_words_in_order,
    but every line ending is a bare LF: a file that has passed through a text
    editor on Linux or macOS must unfold exactly the same way."""
    raw = (
        b"From: Priya <priya@example.com>\n"
        b"To: Ruth <ruth@example.com>,\n"
        b"\tTomas <tomas@example.com>,\n"
        b" Corinne <corinne@example.com>\n"
        b"Subject: Desk coverage\n"
        b"\n"
        b"Body text.\n"
    )
    to_line = next(line for line in _header(raw).split("\n") if line.startswith("To:"))
    assert "\r" not in to_line and "\n" not in to_line
    assert to_line.index("Ruth") < to_line.index("Tomas") < to_line.index("Corinne")


def test_lf_folded_and_crlf_folded_headers_are_byte_identical():
    """Issue #53: the same message, folded once with CRLF and once with bare
    LF, must produce byte-identical header paragraphs."""
    raw_crlf = (
        b"From: a@example.com\r\n"
        b"To: b@example.com\r\n"
        b"Subject: Desk coverage for the week,\r\n"
        b" continuing on the next line\r\n"
        b"\r\n"
        b"Body.\r\n"
    )
    raw_lf = raw_crlf.replace(b"\r\n", b"\n")
    assert _header(raw_crlf) == _header(raw_lf)


@pytest.mark.parametrize(
    "encoded_word",
    ["=?iso-8859-1?Q?caf=E9?=", "=?utf-8?B?Y2Fmw6k=?="],
    ids=["iso-8859-1-Q", "utf-8-B"],
)
def test_an_encoded_word_subject_decodes_the_same_way_for_q_and_b(encoded_word):
    raw = (
        f"From: a@example.com\r\nTo: b@example.com\r\nSubject: {encoded_word}\r\n\r\nBody.\r\n"
    ).encode("ascii")
    result = normalize(raw, "mail/x.eml")
    assert isinstance(result, Conversion)
    assert result.title == "café"
    assert "Subject: café" in result.paragraphs[0]


def test_raw_8bit_from_bytes_decode_to_the_character_with_no_replacement():
    raw = "From: café <a@example.com>\r\nTo: b@example.com\r\nSubject: x\r\n\r\nBody.\r\n".encode("latin-1")
    header = _header(raw)
    assert "café" in header
    assert "�" not in header


def test_no_west_desk_mail_fixture_has_a_cc_line_and_received_only_has_no_date_line():
    for path in sorted(MAIL.glob("*.eml")):
        header = _convert(f"mail/{path.name}").paragraphs[0]
        assert not any(line.startswith("Cc:") for line in header.split("\n")), path.name
    received_only = _convert("mail/received-only.eml").paragraphs[0]
    assert not any(line.startswith("Date:") for line in received_only.split("\n"))


def test_cc_and_date_appear_in_order_when_both_are_present():
    raw = (
        b"From: a@example.com\r\nTo: b@example.com\r\nCc: c@example.com\r\n"
        b"Date: Wed, 23 May 2001 08:58:33 -0500\r\nSubject: x\r\n\r\nBody.\r\n"
    )
    lines = _header(raw).split("\n")
    names = [line.split(":")[0] for line in lines]
    assert names == ["From", "To", "Cc", "Date", "Subject"]


# --- the shared text decode ---------------------------------------------------


def test_only_one_decode_function_is_defined_and_dating_imports_it():
    import strata.dating as dating_module

    assert not hasattr(dating_module, "_decode")
    assert dating_module.decode_text is decode_text


def test_the_shared_decode_handles_cp1252_and_never_raises_on_any_byte():
    result = normalize(b"\x93quoted\x94 caf\xe9", "memos/x.txt")
    assert isinstance(result, Conversion)
    assert result.paragraphs == ("“quoted” café",)
    assert decode_text(b"\x93quoted\x94 caf\xe9") == "“quoted” café"

    for byte in range(256):
        decode_text(bytes([byte]) * 3)  # must not raise


# --- the one suffix map --------------------------------------------------------


def test_converter_ids_names_exactly_the_suffixes_normalize_accepts():
    from strata.normalizer import CONVERTER_IDS

    assert CONVERTER_IDS == {".txt": "text", ".md": "text", ".docx": "docx", ".pdf": "pdf", ".eml": "eml"}
    for suffix in CONVERTER_IDS:
        result = normalize(b"", f"x{suffix}")
        if isinstance(result, Refusal):
            assert "no converter claims" not in result.reason

    unclaimed = normalize(b"whatever", "sources/export.mbox")
    assert isinstance(unclaimed, Refusal)
    assert unclaimed.reason == "no converter claims the suffix '.mbox'"


# --- issue #45: a hand-bumped version travels in the converter id ----------


def test_converter_id_renders_name_at_version_for_every_suffix():
    from strata.normalizer import CONVERTER_IDS, CONVERTER_VERSIONS, converter_id

    for suffix, name in CONVERTER_IDS.items():
        assert converter_id(suffix) == f"{name}@{CONVERTER_VERSIONS[name]}"
    assert converter_id(".mbox") is None


# --- the binary-as-text guard ------------------------------------------------


def test_not_really_text_is_refused_with_a_reason():
    result = normalize((SOURCES / "memos" / "not-really-text.txt").read_bytes(), "memos/not-really-text.txt")
    assert isinstance(result, Refusal)
    assert result.reason


def test_a_long_paragraph_with_almost_no_whitespace_is_refused():
    raw = ("x" * 5000).encode("utf-8")
    result = normalize(raw, "memos/dense.txt")
    assert isinstance(result, Refusal)


def test_ordinary_long_prose_is_not_refused():
    sentence = "The congestion working group met to review the hourly figures for the tie. "
    raw = (sentence * 80).encode("utf-8")
    result = normalize(raw, "memos/long-prose.txt")
    assert isinstance(result, Conversion)


# --- docx and pdf -------------------------------------------------------------


def test_docx_fixtures_convert_with_the_first_paragraph_as_title():
    for name in ("desk-procedures", "outage-letter"):
        conversion = _convert(f"memos/{name}.docx")
        assert conversion.paragraphs
        assert conversion.title == conversion.paragraphs[0]


def test_outage_notice_pdf_converts_to_its_text():
    conversion = _convert("memos/outage-notice.pdf")
    assert "Planned outage, Cascade tie" in conversion.paragraphs[0]


def test_signed_statement_scan_pdf_has_no_text_layer_and_a_reported_reason():
    result = normalize(
        (SOURCES / "memos" / "signed-statement-scan.pdf").read_bytes(), "memos/signed-statement-scan.pdf"
    )
    assert isinstance(result, Conversion)
    assert result.paragraphs == ()
    assert result.metadata["reason"]


# --- unclaimed suffixes and markdown headings --------------------------------


def test_an_unclaimed_suffix_is_refused():
    result = normalize(b"whatever", "sources/export.mbox")
    assert isinstance(result, Refusal)


def test_a_markdown_heading_becomes_the_title():
    result = normalize(b"# The Outage\n\nBody text here.", "notes/theme/outage.md")
    assert isinstance(result, Conversion)
    assert result.title == "The Outage"
    assert result.paragraphs == ("# The Outage", "Body text here.")


# --- byte-stable and date-free -------------------------------------------------


ALL_FIXTURES = (
    [f"mail/{p.name}" for p in sorted(MAIL.glob("*.eml"))]
    + ["memos/desk-procedures.docx", "memos/outage-letter.docx"]
    + ["memos/outage-notice.pdf", "memos/signed-statement-scan.pdf"]
    + ["memos/glossary.txt", "2001/May/desk-checklist.txt"]
)


@pytest.mark.parametrize("relative", ALL_FIXTURES)
def test_converter_output_is_byte_stable_across_two_runs(relative):
    raw = (SOURCES / relative).read_bytes()
    first = normalize(raw, relative)
    second = normalize(raw, relative)
    assert first == second


@pytest.mark.parametrize("relative", ALL_FIXTURES)
def test_no_converter_sets_or_reads_a_date(relative):
    raw = (SOURCES / relative).read_bytes()
    result = normalize(raw, relative)
    if isinstance(result, Conversion):
        assert not any("date" in key.lower() for key in result.metadata)
        for value in result.metadata.values():
            assert isinstance(value, str)
            assert not _ISO_DAY.search(value), f"{relative}: metadata carries a date: {value!r}"
        assert not _ISO_DAY.search(result.title), f"{relative}: title carries an ISO date: {result.title!r}"


_ISO_DAY = re.compile(r"\b(19|20)\d\d-\d\d-\d\d\b")


def test_an_email_date_header_is_copied_verbatim_and_read_by_nothing_else():
    """Changing only the Date header changes only the header paragraph's
    Date line, and that line is the header's own wording: nothing parses,
    reformats or derives anything from it."""
    raw = (MAIL / "plain.eml").read_bytes()
    header_line = next(line for line in raw.splitlines() if line.startswith(b"Date:"))
    replaced = raw.replace(header_line, b"Date: Sun, 7 Jan 1990 03:04:05 +0900", 1)

    original = normalize(raw, "mail/plain.eml")
    changed = normalize(replaced, "mail/plain.eml")
    assert isinstance(original, Conversion) and isinstance(changed, Conversion)

    date_line = next(line for line in changed.paragraphs[0].split("\n") if line.startswith("Date:"))
    assert date_line == "Date: Sun, 7 Jan 1990 03:04:05 +0900"
    assert changed.paragraphs[0].replace(date_line, "") == original.paragraphs[0].replace(
        next(line for line in original.paragraphs[0].split("\n") if line.startswith("Date:")), ""
    )
    assert changed.paragraphs[1:] == original.paragraphs[1:]
    assert changed.title == original.title
    assert changed.metadata == original.metadata


def test_a_docx_created_property_does_not_reach_the_conversion(tmp_path):
    import docx

    def build(created: datetime) -> bytes:
        document = docx.Document()
        document.add_paragraph("Desk rota for the week.")
        document.core_properties.created = created
        path = tmp_path / f"rota-{created.year}.docx"
        document.save(str(path))
        return path.read_bytes()

    early = normalize(build(datetime(1990, 1, 7, 3, 4, 5)), "memos/rota.docx")
    late = normalize(build(datetime(2001, 6, 19, 23, 10, 0)), "memos/rota.docx")
    assert isinstance(early, Conversion)
    assert early == late
    assert early.metadata == {}


# --- docx heading styles ---------------------------------------------------------


def _docx_with_heading(tmp_path, *, level: int) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Circulated to the west desk only.")
    document.add_heading("Revised cutoff procedure", level=level)
    document.add_paragraph("Schedules after the cutoff go to Ruth first.")
    path = tmp_path / f"heading-{level}.docx"
    document.save(str(path))
    return path.read_bytes()


@pytest.mark.parametrize("level", [0, 1, 2])
def test_a_docx_heading_or_title_style_paragraph_becomes_the_title(tmp_path, level):
    """Level 0 is python-docx's ``Title`` style, 1 and 2 are ``Heading 1``
    and ``Heading 2``. The first such paragraph is the title even when a
    Normal paragraph precedes it, and it stays in the paragraphs as text."""
    result = normalize(_docx_with_heading(tmp_path, level=level), "memos/cutoff.docx")
    assert isinstance(result, Conversion)
    assert result.title == "Revised cutoff procedure"
    assert result.paragraphs == (
        "Circulated to the west desk only.",
        "Revised cutoff procedure",
        "Schedules after the cutoff go to Ruth first.",
    )


def test_only_the_first_docx_heading_becomes_the_title(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("First heading", level=1)
    document.add_paragraph("Body.")
    document.add_heading("Second heading", level=1)
    path = tmp_path / "two-headings.docx"
    document.save(str(path))
    result = normalize(path.read_bytes(), "memos/two-headings.docx")
    assert isinstance(result, Conversion)
    assert result.title == "First heading"
